from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'report'/'LAB03_House_Price_BaoCao.docx'

def set_cell_shading(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def add_table(doc, headers, rows):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    hdr=table.rows[0]
    set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        hdr.cells[i].text=h; set_cell_shading(hdr.cells[i]); hdr.cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in hdr.cells[i].paragraphs[0].runs: r.bold=True
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table

def p(doc, text='', bold=False, italic=False, align=None, size=None):
    par=doc.add_paragraph();
    if align is not None: par.alignment=align
    run=par.add_run(text); run.bold=bold; run.italic=italic
    if size: run.font.size=Pt(size)
    return par

def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.9); sec.right_margin=Inches(0.9)
styles=doc.styles
styles['Normal'].font.name='Times New Roman'; styles['Normal'].font.size=Pt(12)
styles['Title'].font.name='Times New Roman'; styles['Title'].font.size=Pt(22)
for s in ['Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Times New Roman'

# Cover
p(doc,'TRƯỜNG ĐẠI HỌC ........................................',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=15)
p(doc,'KHOA / BỘ MÔN ........................................',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=13)
p(doc,'',size=12)
p(doc,'BÁO CÁO LAB 03',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=20)
p(doc,'HOUSE PRICES - ADVANCED REGRESSION TECHNIQUES',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=17)
p(doc,'',size=18)
p(doc,'Đề tài: Dự đoán giá nhà bằng Machine Learning và MLP',align=WD_ALIGN_PARAGRAPH.CENTER,size=14)
p(doc,'',size=18)
add_table(doc,['Thông tin','Nội dung'],[
    ['Môn học','Machine Learning / Data Science'],
    ['Nhóm trưởng','........................................................'],
    ['Mã SV','........................................................'],
    ['Thành viên khác','........................................................'],
    ['Tên file nộp','lab03_house_price_hoten_masv.zip'],
])
p(doc,'',size=20)
p(doc,'Năm học: 2026 - 2027',align=WD_ALIGN_PARAGRAPH.CENTER,size=12)
doc.add_page_break()

# Contents
heading(doc,'MỤC LỤC',1)
for x in [
    '1. Giới thiệu bài toán', '2. Cơ sở dữ liệu và paper De Cock', '3. Phân tích dữ liệu',
    '4. Tiền xử lý và xây dựng đặc trưng', '5. Mô hình Machine Learning bằng scikit-learn',
    '6. Mô hình MLP bằng PyTorch', '7. Thực nghiệm và kiểm tra', '8. Kết quả và nhận xét',
    '9. Kết luận', '10. Tài liệu tham khảo']:
    p(doc,x)
doc.add_page_break()

heading(doc,'PHÂN CÔNG CÔNG VIỆC',1)
add_table(doc,['Thành viên','Công việc'],[
    ['................................','Đọc paper, tìm hiểu dữ liệu Ames Housing và xác định bài toán hồi quy.'],
    ['................................','Tiền xử lý dữ liệu, xử lý missing values và one-hot encoding.'],
    ['................................','Xây dựng mô hình scikit-learn và thực nghiệm feature engineering.'],
    ['................................','Xây dựng MLP bằng PyTorch, kiểm tra và tạo file submission.'],
    ['Nhóm trưởng','Tổng hợp source code, báo cáo và nén file nộp.'],
])

heading(doc,'1. GIỚI THIỆU BÀI TOÁN',1)
p(doc,'Bài toán của Lab 03 là xây dựng mô hình dự đoán giá bán của một căn nhà dựa trên các thuộc tính mô tả căn nhà. Dữ liệu được sử dụng là bộ House Prices - Advanced Regression Techniques trên Kaggle, có 1.460 mẫu train, 1.459 mẫu test và 79 biến giải thích. Biến mục tiêu là SalePrice.')
p(doc,'Đây là bài toán hồi quy (regression) vì đầu ra là một giá trị số liên tục. Mục tiêu của mô hình là học quan hệ giữa các đặc trưng của ngôi nhà và giá bán, sau đó dự đoán SalePrice cho tập test.')
p(doc,'Kaggle đánh giá bằng RMSE trên logarit của giá dự đoán và giá thực tế. Vì vậy trong quá trình huấn luyện, báo cáo sử dụng y_log = log1p(SalePrice) và tính RMSLE trên tập validation để gần với tiêu chí của cuộc thi.')
p(doc,'Nguồn Kaggle mô tả cuộc thi là bài tập dự đoán SalePrice với 79 biến giải thích, đồng thời nhấn mạnh feature engineering và các phương pháp hồi quy / ensemble.')

heading(doc,'2. CƠ SỞ DỮ LIỆU VÀ PAPER DE COCK',1)
p(doc,'Paper “Ames, Iowa: Alternative to the Boston Housing Data as an End of Semester Regression Project” của Dean De Cock trình bày dữ liệu bán nhà tại Ames, Iowa trong giai đoạn 2006-2010. Phiên bản dữ liệu được giới thiệu trong paper có 2.930 quan sát sau khi loại các giao dịch không phải nhà ở và giữ lại lần bán gần nhất của các căn bị lặp.')
p(doc,'Paper cho biết dữ liệu có nhiều kiểu biến: 23 biến nominal, 23 biến ordinal, 14 biến discrete và 20 biến continuous. Các biến mô tả nhiều thuộc tính vật lý của căn nhà như diện tích, phòng ngủ, phòng tắm, tầng hầm, garage và khu vực.')
p(doc,'Một điểm quan trọng trong paper là xử lý mô hình theo hai giai đoạn: mô hình đơn giản và mô hình phức tạp. Paper cũng khuyến nghị chia dữ liệu thành train và validation, kiểm tra mô hình trên validation và thực hiện một phép kiểm tra riêng đối với quan sát đầu tiên để phát hiện lỗi khi tạo biến.')
p(doc,'Paper cũng lưu ý đến outlier. Tác giả đề nghị loại các căn có diện tích Gr Liv Area trên 4.000 square feet trong bối cảnh bài tập của mình trước khi giao dữ liệu cho sinh viên. Trong Lab 03, bước này được xem như một quy tắc tham khảo; source code chính không tự động xóa dữ liệu để tránh làm thay đổi tập Kaggle mà không có sự kiểm tra của người dùng.')

heading(doc,'3. PHÂN TÍCH DỮ LIỆU',1)
p(doc,'Từ code tham khảo được cung cấp, quy trình ban đầu gồm đọc train.csv và test.csv, kiểm tra shape, kiểu dữ liệu và số lượng giá trị thiếu. Code tham khảo cũng kiểm tra missing bằng heatmap, loại một số cột có tỷ lệ thiếu rất cao, điền mode cho biến phân loại và mean cho một số biến số, sau đó one-hot encoding.')
add_table(doc,['Nội dung','Kết quả / cách xử lý'],[
    ['Train','1.460 dòng, 81 cột (gồm Id và SalePrice).'],
    ['Test','1.459 dòng, 80 cột (gồm Id, không có SalePrice).'],
    ['Target','SalePrice'],
    ['Biến phân loại','One-hot encoding'],
    ['Biến số','Median imputation; StandardScaler cho MLP'],
    ['Cột thưa','Xem tỷ lệ thiếu trên train; các cột vượt ngưỡng được loại.'],
    ['Id','Loại khỏi đặc trưng đưa vào mô hình.'],
])

heading(doc,'4. TIỀN XỬ LÝ VÀ XÂY DỰNG ĐẶC TRƯNG',1)
p(doc,'Source code sử dụng một pipeline chung để giảm thao tác thủ công và bảo đảm cách biến đổi dữ liệu train/validation/test nhất quán. Các biến phân loại được điền giá trị thiếu bằng mode và OneHotEncoder(handle_unknown="ignore"). Các biến số được điền median; riêng MLP được chuẩn hóa bằng StandardScaler.')
heading(doc,'4.1. Các đặc trưng tổng hợp',2)
add_table(doc,['Đặc trưng','Cách tạo','Ý nghĩa'],[
    ['TotalSF','TotalBsmtSF + 1stFlrSF + 2ndFlrSF','Tổng diện tích sử dụng.'],
    ['TotalBathrooms','FullBath + 0.5*HalfBath + BsmtFullBath + 0.5*BsmtHalfBath','Quy đổi số phòng tắm về một thang chung.'],
    ['TotalPorchSF','WoodDeckSF + OpenPorchSF + EnclosedPorch + 3SsnPorch + ScreenPorch','Tổng diện tích khu vực hiên / deck.'],
    ['HouseAge','YrSold - YearBuilt','Tuổi căn nhà tại thời điểm bán.'],
    ['RemodAge','YrSold - YearRemodAdd','Số năm từ lần cải tạo gần nhất.'],
    ['GarageAge','YrSold - GarageYrBlt','Tuổi garage tại thời điểm bán.'],
    ['Total_sqr_footage','BsmtFinSF1 + BsmtFinSF2 + 1stFlrSF + 2ndFlrSF','Diện tích tính theo các phần chính.'],
    ['OverallQual_GrLivArea','OverallQual * GrLivArea','Kết hợp chất lượng tổng thể và diện tích ở.'],
    ['OverallQual_TotalSF','OverallQual * TotalSF','Kết hợp chất lượng tổng thể và tổng diện tích.'],
])
heading(doc,'4.2. Log-transform',2)
p(doc,'Với biến số không âm có độ lệch lớn, source code có một thực nghiệm riêng sử dụng log1p. Mục tiêu là làm giảm ảnh hưởng của phân phối lệch phải. Cách này được áp dụng ở biến thể engineered_log và đặc biệt hữu ích cho MLP vì mạng nơ-ron nhạy với scale của đầu vào.')

heading(doc,'5. MÔ HÌNH MACHINE LEARNING BẰNG SCIKIT-LEARN',1)
p(doc,'Mô hình truyền thống được chọn là RandomForestRegressor. Đây là mô hình ensemble gồm nhiều cây quyết định; mỗi cây được huấn luyện trên một mẫu bootstrap và tại mỗi nút chỉ xét một tập con đặc trưng. Mô hình phù hợp với dữ liệu nhiều biến và có khả năng học quan hệ phi tuyến.')
p(doc,'Trong source code, RandomForestRegressor được dùng trên dữ liệu đã qua preprocessing. Target được biến đổi thành log1p(SalePrice). Sau khi dự đoán, dùng expm1 để đưa giá dự đoán về đơn vị dollar.')
add_table(doc,['Tham số','Giá trị'],[
    ['n_estimators','500'],['max_features','0.7'],['min_samples_leaf','1'],['random_state','42'],['n_jobs','-1']])

heading(doc,'6. MÔ HÌNH MLP BẰNG PYTORCH',1)
p(doc,'Mô hình MLP (Multi-Layer Perceptron) được xây dựng bằng PyTorch để thay cho phần Artificial Neural Network sử dụng Keras trong code tham khảo. Kiến trúc được giữ đơn giản, phù hợp với bài lab: lớp đầu vào -> Dense 128 -> ReLU -> Dropout -> Dense 64 -> ReLU -> Dropout -> đầu ra 1 giá trị.')
add_table(doc,['Thành phần','Thiết lập'],[
    ['Input','Số chiều bằng số cột sau one-hot encoding'],
    ['Hidden layer 1','128 neurons, ReLU'],
    ['Dropout','0.20'],
    ['Hidden layer 2','64 neurons, ReLU'],
    ['Dropout','0.10'],
    ['Output','1 neuron'],
    ['Loss','MSE trên log1p(SalePrice)'],
    ['Optimizer','Adam'],
    ['Learning rate','0.001'],
    ['Early stopping','Patience 35'],
])

heading(doc,'7. THỰC NGHIỆM VÀ KIỂM TRA',1)
p(doc,'Source code chạy 4 biến thể đặc trưng cho Random Forest: BASELINE, ENGINEERED, ENGINEERED_LOG và SELECTED. Sau đó chạy MLP trên ENGINEERED_LOG. Cách chia validation là 80/20 với random_state = 42, tương tự tinh thần split train/validation được trình bày trong paper.')
add_table(doc,['Thực nghiệm','Mục đích'],[
    ['BASELINE','Làm mốc ban đầu với đặc trưng gốc sau preprocessing.'],
    ['ENGINEERED','Kiểm tra tác động của các biến tổng hợp và biến tuổi.'],
    ['ENGINEERED_LOG','Kiểm tra thêm log1p trên các biến số lệch phải.'],
    ['SELECTED','Kiểm tra một tập đặc trưng gọn, dễ giải thích.'],
    ['PyTorch MLP','Kiểm tra khả năng học phi tuyến của mạng nơ-ron.'],
])
heading(doc,'7.1. Kiểm tra source code',2)
p(doc,'File src/smoke_test.py cung cấp một bài kiểm tra nhanh với dữ liệu giả lập nhỏ. Bài test kiểm tra việc tạo đặc trưng, OneHotEncoder, StandardScaler và bảo đảm ma trận đầu vào không chứa NaN/Inf sau preprocessing. Bài test đã được chạy thành công trong quá trình tạo gói: SMOKE TEST PASSED; 30 cột đầu vào và 49 cột sau biến đổi.')
p(doc,'Khi có train.csv và test.csv thật, chạy python src/run_all.py. Script sẽ tạo results_summary.csv, learning curve của MLP, top 20 feature importance, manual_check.txt và hai file submission cho Kaggle.')

heading(doc,'7.2. Kiểm tra theo tinh thần paper',2)
p(doc,'Paper nhấn mạnh việc không nên phụ thuộc hoàn toàn vào output phần mềm và yêu cầu kiểm tra lại một quan sát bằng phép tính độc lập. Source code có tạo manual_check.txt cho quan sát đầu tiên nhằm rà lỗi pipeline. Khi nộp báo cáo, nếu giảng viên yêu cầu phép tính tay hoàn toàn, cần lấy đúng các giá trị gốc của một quan sát, tính các biến biến đổi và thay vào phương trình / pipeline tương ứng.')

heading(doc,'8. KẾT QUẢ VÀ NHẬN XÉT',1)
p(doc,'Do các file train.csv và test.csv của Kaggle không nằm trong hai file PDF được đính kèm của đề, gói source không ghi số RMSLE giả. Đây là lựa chọn để bảo đảm tính trung thực của báo cáo. Sau khi đặt dữ liệu thật vào thư mục data và chạy src/run_all.py, bảng results/results_summary.csv sẽ chứa số RMSLE validation của từng thực nghiệm.')
add_table(doc,['Model / Variant','RMSLE validation','Transformed features','Ghi chú'],[
    ['RandomForest - BASELINE','................','................','Mốc ban đầu'],
    ['RandomForest - ENGINEERED','................','................','Thêm đặc trưng tổng hợp'],
    ['RandomForest - ENGINEERED_LOG','................','................','Thêm log1p'],
    ['RandomForest - SELECTED','................','................','Tập đặc trưng rút gọn'],
    ['PyTorch MLP - ENGINEERED_LOG','................','................','Mạng nơ-ron'],
])
p(doc,'Cách nhận xét trong báo cáo: so sánh RMSLE của từng biến thể với BASELINE; nếu RMSLE giảm thì biến thể giúp lỗi dự đoán trên tập validation giảm. Không cần ghi “tốt nhất” theo cảm tính; nên mô tả trực tiếp giá trị đo được và sự thay đổi so với mốc ban đầu.')

heading(doc,'9. KẾT LUẬN',1)
p(doc,'Lab 03 đã xây dựng một quy trình hoàn chỉnh từ dữ liệu House Prices đến preprocessing, feature engineering, mô hình Random Forest bằng scikit-learn và MLP bằng PyTorch. Thiết kế được bám theo tinh thần paper De Cock: chú ý kiểu dữ liệu, missing values, feature construction, train/validation split và kiểm tra kết quả bằng một quan sát riêng.')
p(doc,'Điểm có thể tiếp tục mở rộng là thử thêm GradientBoostingRegressor, ElasticNet hoặc các ensemble khác trong phạm vi môn học; tuy nhiên source nộp chính của Lab 03 giữ hai mô hình đúng yêu cầu là scikit-learn và PyTorch.')

heading(doc,'10. TÀI LIỆU THAM KHẢO',1)
p(doc,'[1] Dean De Cock (2011), “Ames, Iowa: Alternative to the Boston Housing Data as an End of Semester Regression Project”, Journal of Statistics Education, Volume 19, Number 3.')
p(doc,'[2] Kaggle, “House Prices - Advanced Regression Techniques”, Competition Overview / Data / Evaluation.')
p(doc,'[3] File code tham khảo được cung cấp trong đề: ML_House Prices_Advanced Regression Techniques.pdf.')
p(doc,'[4] Scikit-learn documentation: RandomForestRegressor, ColumnTransformer, OneHotEncoder, StandardScaler.')
p(doc,'[5] PyTorch documentation: torch.nn, Adam optimizer, model training.')

# Footer page number not dynamic; simple footer text
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('LAB03 - HOUSE PRICE').font.size=Pt(9)

doc.save(OUT)
print(OUT)
